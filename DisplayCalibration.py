"""Shared access to the calibrated display gray levels used by tasks."""

import json
import math
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_BASE_PATH = SCRIPT_DIR / "stim_calibration_0_20.json"
DEFAULT_EXTENSION_PATH = SCRIPT_DIR / "stim_calibration_extension_to_255.json"
DEFAULT_EXTENDED_WORKBOOK_PATH = (
    SCRIPT_DIR / "屏幕RGB_扩展等亮度刺激标定_至255.xlsx"
)


def _read_json(path, description):
    path = Path(path)
    try:
        with path.open("r", encoding="utf-8-sig") as source_file:
            payload = json.load(source_file)
    except (OSError, ValueError) as exc:
        raise RuntimeError(f"Cannot read {description}: {path}") from exc
    if not isinstance(payload, dict):
        raise RuntimeError(f"{description} must contain a JSON object: {path}")
    return payload


def _inverse_interpolate_luminance(target_luminance, curve):
    if target_luminance <= curve[0][1]:
        return curve[0][0]
    if target_luminance >= curve[-1][1]:
        return curve[-1][0]
    for index in range(1, len(curve)):
        lower_input, lower_luminance = curve[index - 1]
        upper_input, upper_luminance = curve[index]
        if target_luminance <= upper_luminance:
            luminance_span = upper_luminance - lower_luminance
            if luminance_span <= 0:
                return upper_input
            fraction = (
                (target_luminance - lower_luminance) / luminance_span
            )
            return lower_input + fraction * (upper_input - lower_input)
    return curve[-1][0]


def load_calibrated_gray_level(
    level,
    base_path=DEFAULT_BASE_PATH,
    extension_path=DEFAULT_EXTENSION_PATH,
):
    """Return RGB and luminance for one calibrated Gray intensity level."""
    if isinstance(level, bool) or int(level) != float(level):
        raise ValueError("Background Gray Level must be an integer.")
    level = int(level)

    base = _read_json(base_path, "base display calibration")
    extension = _read_json(extension_path, "extended display calibration")
    base_luminance = base.get("target_luminance_cd_m2", [])
    base_gray = base.get("rgb_by_color", {}).get("Gray", [])
    if len(base_luminance) != 21 or len(base_gray) != 21:
        raise ValueError(
            "Base display calibration must contain Gray levels 0-20."
        )

    max_level = int(extension["max_level_by_color"]["Gray"])
    if not 0 <= level <= max_level:
        raise ValueError(
            f"Background Gray Level must be between 0 and {max_level}."
        )

    if level <= 20:
        rgb = tuple(int(channel) for channel in base_gray[level])
        luminance = float(base_luminance[level])
    else:
        curve = [
            (float(input_value), float(luminance_value))
            for input_value, luminance_value in extension[
                "measured_curve_by_color"
            ]["Gray"]
        ]
        target_luminance = min(
            level * float(extension["luminance_step_cd_m2"]),
            curve[-1][1],
        )
        input_value = _inverse_interpolate_luminance(
            target_luminance, curve
        )
        integer_input = min(
            255,
            max(0, int(math.floor(input_value + 0.5))),
        )
        rgb = (integer_input, integer_input, integer_input)
        luminance = float(target_luminance)

    return {
        "level": level,
        "rgb_255": rgb,
        "luminance_cd_m2": luminance,
        "max_level": max_level,
        "base_path": str(Path(base_path).resolve()),
        "extension_path": str(Path(extension_path).resolve()),
    }


def load_calibrated_gray_level_from_workbook(
    level,
    workbook_path=DEFAULT_EXTENDED_WORKBOOK_PATH,
):
    """Read one Gray brightness level directly from the extended workbook."""
    if isinstance(level, bool) or int(level) != float(level):
        raise ValueError("Gray brightness level must be an integer.")
    level = int(level)
    workbook_path = Path(workbook_path)

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to read the display calibration workbook."
        ) from exc

    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=True,
        )
    except (OSError, ValueError) as exc:
        raise RuntimeError(
            f"Cannot read extended display calibration: {workbook_path}"
        ) from exc

    try:
        if "Gray" not in workbook.sheetnames:
            raise ValueError(
                "Extended display calibration must contain a Gray sheet."
            )
        sheet = workbook["Gray"]
        for row in sheet.iter_rows(
            min_row=7,
            max_col=8,
            values_only=True,
        ):
            row_level = row[0]
            if row_level is None:
                continue
            if int(row_level) != level:
                continue

            rgb = tuple(int(row[index]) for index in (4, 5, 6))
            return {
                "level": level,
                "rgb_255": rgb,
                "luminance_cd_m2": float(row[1]),
                "estimated_luminance_cd_m2": float(row[7]),
                "max_level": int(sheet["F3"].value),
                "workbook_path": str(workbook_path.resolve()),
                "sheet": "Gray",
            }
    finally:
        workbook.close()

    raise ValueError(
        f"Gray brightness level must be between 0 and the workbook maximum: "
        f"{level} was not found."
    )


def rgb255_to_psychopy(rgb):
    return [int(channel) / 127.5 - 1.0 for channel in rgb]
